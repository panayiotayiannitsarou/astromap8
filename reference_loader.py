import io
from pathlib import Path

import streamlit as st
from docx import Document


REFERENCE_DIR = Path(__file__).resolve().parent / "references"
DEFAULT_INSTRUCTIONS = REFERENCE_DIR / "Odigies_v5.docx"
DEFAULT_STYLE = REFERENCE_DIR / "Elena_style_guide_v2.docx"
ROOT_INSTRUCTIONS = Path(__file__).resolve().parent / "Odigies_v5.docx"
ROOT_STYLE = Path(__file__).resolve().parent / "Elena_style_guide_v2.docx"
COMMON_ORIENTATION = Path(__file__).resolve().parent / "Desmeftiki_Entoli_Epaggelmatikou_Prosanatolismou_Koini_v3.docx"


def docx_text(source) -> str:
    """Extract paragraphs and tables from a DOCX path or uploaded bytes."""
    if isinstance(source, (bytes, bytearray)):
        source = io.BytesIO(source)
    document = Document(source)
    blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            line = " | ".join(cell.text.strip() for cell in row.cells)
            if line.strip(" |"):
                blocks.append(line)
    return "\n".join(blocks)


@st.cache_data(show_spinner=False)
def load_default_references() -> tuple[str, str]:
    # @st.cache_data: το Streamlit ξανατρέχει ολόκληρο το script σε κάθε
    # interaction, οπότε χωρίς caching αυτά τα (συχνά εκτενή) .docx
    # ξαναδιαβάζονταν και ξαναπαρσάρονταν από τον δίσκο σε κάθε κλικ.
    #
    # Accept both repository layouts: a dedicated references/ folder or the
    # two DOCX files beside app.py.  This makes GitHub web uploads simpler.
    instructions = DEFAULT_INSTRUCTIONS if DEFAULT_INSTRUCTIONS.exists() else ROOT_INSTRUCTIONS
    style = DEFAULT_STYLE if DEFAULT_STYLE.exists() else ROOT_STYLE
    if not instructions.exists() or not style.exists():
        # Πριν έγραφε "...v4...", ενώ το πραγματικό αρχείο είναι Odigies_v5.docx
        # (και το app.py το παρουσιάζει ως "Ενσωματωμένες οδηγίες v5.3") --
        # ένα μήνυμα σφάλματος έπρεπε τουλάχιστον να συμφωνεί με το filename.
        raise FileNotFoundError("Λείπουν οι ενσωματωμένες οδηγίες v5 ή το πρότυπο ύφους.")
    return docx_text(instructions), docx_text(style)


def load_orientation_command(service: str) -> str:
    if not COMMON_ORIENTATION.exists():
        raise FileNotFoundError(f"Λείπει η κοινή εντολή προσανατολισμού: {COMMON_ORIENTATION.name}")
    return docx_text(COMMON_ORIENTATION)
